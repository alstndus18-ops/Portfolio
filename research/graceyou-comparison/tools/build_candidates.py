"""후보 풀(data/candidates.json) + 플랫폼 스냅샷(snapshots/) + 기사 검색 결과(data/raw/news)를 합쳐
  candidates/<slug>.md · candidates/<slug>.json (브랜드별 파일), candidates.csv · candidates.xlsx (전체 표), candidates/INDEX.md 를 만든다.
  python tools/build_candidates.py
점수는 tools/build.py 와 같은 3축이다: 가격 위치(0.30) · 가격대 겹침(0.20) · 무드 겹침(0.50). 상품 수·카테고리는 점수에 쓰지 않는다.
가격이나 무드 태그가 없는 후보는 점수를 비우고 '미평가'로 둔다(후보에서 빼지 않는다).
"""
import csv
import glob
import hashlib
import json
import math
import re
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
D = ROOT / "data"
S = ROOT / "snapshots"
W = {"price": 0.30, "band": 0.20, "mood": 0.50}

reg = json.loads((D / "candidates.json").read_text(encoding="utf-8"))
meta = json.loads((D / "brand_meta.json").read_text(encoding="utf-8"))
GU = meta["graceu"]
GP = GU["price_uniq_all"]
GT = set(GU["tags"])
off = json.loads((D / "tops_summary.json").read_text(encoding="utf-8"))["brands"]  # tools/build_onepager.py 가 만든 상의 요약


def latest(kind, key):
    fs = sorted(glob.glob(str(S / kind / str(key) / "*.json")))
    return json.loads(Path(fs[-1]).read_text(encoding="utf-8")) if fs else None


def dates(kind, key):
    return [Path(f).stem for f in sorted(glob.glob(str(S / kind / str(key) / "*.json")))]


def slugify(b):
    en = (b.get("en") or "").lower()
    s = re.sub(r"[^a-z0-9]+", "-", en).strip("-")
    if not s:
        ids = b.get("ids", {})
        s = "b" + str(ids.get("c29") or ids.get("wc") or hashlib.md5(b["ko"].encode("utf-8")).hexdigest()[:6])  # 실행마다 같은 이름이 되도록 고정 해시
    return s


# ---------- 기사 근거 ----------
news = {}
for f in sorted(glob.glob(str(D / "raw" / "news" / "*_apparelnews.json"))):
    for kw, items in json.loads(Path(f).read_text(encoding="utf-8")).items():
        for it in items:
            news.setdefault(it["idx"], dict(it, kws=set()))["kws"].add(kw)
ctx = {}
for f in sorted(glob.glob(str(D / "raw" / "news" / "context_*.json"))):
    for anchor, arts in json.loads(Path(f).read_text(encoding="utf-8")).items():
        for a in arts:
            for c in a.get("context", []):
                ctx.setdefault(a["idx"], []).append((anchor, c))
ANCHORS = {"그레이스유": "그레이스유", "낫띵리튼": "낫띵리튼", "쿠메": "쿠메"}


def adate(it):
    m = re.search(r"(\d{4}-\d{2}-\d{2})", " ".join(it.get("meta", [])))
    return m.group(1) if m else ""


def tok(name):
    return re.compile(r"(?<![가-힣A-Za-z0-9])" + re.escape(name) + r"(?![가-힣A-Za-z0-9])")


def evidence(b):
    names = [b["ko"]] + ([b["en"]] if b.get("en") and len(b["en"]) >= 4 else [])
    pats = [tok(n) for n in names]
    out = {}
    for idx, it in news.items():
        where = None
        if any(p.search(it["title"]) for p in pats):
            where = "제목"
        elif any(p.search(it.get("snippet", "")) for p in pats):
            where = "요약"
        if where:
            out[idx] = {"idx": idx, "date": adate(it), "title": it["title"], "url": it["url"], "where": where, "anchor": ""}
    for idx, lst in ctx.items():
        for anchor, c in lst:
            if any(p.search(c) for p in pats) and anchor != b["ko"]:
                it = news.get(idx, {})
                e = out.setdefault(idx, {"idx": idx, "date": adate(it), "title": it.get("title", ""), "url": it.get("url", ""), "where": "본문 발췌", "anchor": ""})
                e["anchor"] = anchor
                e["excerpt"] = c[:260]
    pri = {"제목": 0, "본문 발췌": 1, "요약": 2}
    return sorted(out.values(), key=lambda e: (pri[e["where"]], -int(e["date"].replace("-", "") or 0)))


# ---------- 가격·점수 ----------
def est_factor(snap):
    d = snap.get("disc") or {}
    if d.get("n") and d["n"] >= snap["n"] * 0.5 and d.get("med"):
        return 1 / (1 - d["med"] / 100)
    return 1.0


def overlap(a1, a2, b1, b2):
    lo, hi = max(a1, b1), min(a2, b2)
    inter = max(0, math.log(hi) - math.log(lo)) if hi > lo else 0
    union = math.log(max(a2, b2)) - math.log(min(a1, b1))
    return inter / union if union else 0


def score(snap, tags):
    if not snap or not snap.get("p") or not snap["p"].get("med"):
        return {}
    k = est_factor(snap)
    p = snap["p"]
    med, q1, q3 = (max(1, round(p[x] * k)) for x in ("med", "q1", "q3"))
    sp = max(0.0, 1 - abs(math.log(med / GP["med"])) / math.log(2.5))
    sb = overlap(q1, q3, GP["q1"], GP["q3"]) if q3 > q1 else 0.0
    r = {"k": round(k, 2), "list_q1": q1, "list_med": med, "list_q3": q3, "s_price": round(sp * 100), "s_band": round(sb * 100)}
    if tags:
        t = set(tags)
        sm = len(t & GT) / len(t | GT)
        r["s_mood"] = round(sm * 100)
        r["score"] = round((W["price"] * sp + W["band"] * sb + W["mood"] * sm) * 100)
    r["partial"] = round((W["price"] * sp + W["band"] * sb) / (W["price"] + W["band"]) * 100)
    return r


def role(b, sc):
    st = b["status"]
    if st == "확정":
        return "직접 경쟁(확정)"
    if st == "제외-가격대":
        return "저가 참고(상의 후보에서 제외)"
    if st == "추천":
        return "경쟁사 추천(우아함 대체재)"
    if "의류 아님" in b.get("note", "") or "가방 브랜드" in b.get("note", "") or "가방·신발 위주" in b.get("note", ""):
        return "비의류·참고"
    if not sc:
        return "미평가(가격 미수집)"
    med = sc["list_med"]
    ratio = med / GP["med"]
    if ratio < 0.5:
        return "가격 앵커(저가)"
    if ratio > 1.6:
        return "가격 앵커(고가)"
    sm, spr = sc.get("s_mood"), sc["s_price"]
    if sm is None:
        return "가격 근접·무드 미평가"
    if spr >= 70 and sm >= 30:
        return "가격·무드 모두 근접"
    if spr >= 70:
        return "가격 근접·무드 다름"
    if sm >= 40:
        return "무드 근접·가격 다름"
    return "참고(트렌드·무드 탐색)"


# ---------- 합치기 ----------
rows = []
base = reg["baseline"]
bw = latest("wconcept", base["ids"]["wc"])
gu_wc = {"n": bw["n"], "p": bw["p"], "disc": bw["disc"]} if bw else None
for b in reg["brands"]:
    ids = b.get("ids", {})
    slug = slugify(b)
    if any(r["slug"] == slug for r in rows):
        slug += "-" + str(ids.get("c29") or ids.get("wc"))
    s29 = latest("29cm", ids["c29"]) if ids.get("c29") else None
    swc = latest("wconcept", ids["wc"]) if ids.get("wc") else None
    m = meta["brands"].get(str(ids.get("c29"))) if ids.get("c29") else None
    tags = b.get("tags") or (m or {}).get("tags") or []
    tier = b.get("tier") or (m or {}).get("tier") or ""
    src = s29 or swc
    src_name = "29CM" if s29 else ("W컨셉" if swc else "")
    sc = score(src, tags)
    if tags is not None and sc is not None:
        pass
    r = {"slug": slug, "ko": b["ko"], "en": b.get("en") or (s29 or swc or {}).get("en") or "", "status": b["status"], "tier": tier, "tags": tags,
         "found": b.get("found", []), "note": b.get("note", ""), "ids": ids, "price_source": src_name,
         "intro": (s29 or swc or {}).get("intro", ""), "sub": (s29 or {}).get("sub", ""), "est": (s29 or {}).get("est"), "heart": (s29 or {}).get("heart"),
         "s29": s29, "swc": swc, "off": off.get(b["ko"]), "score": sc, "evidence": evidence(b),
         "snap_days": {"29cm": dates("29cm", ids["c29"]) if ids.get("c29") else [], "wconcept": dates("wconcept", ids["wc"]) if ids.get("wc") else []},
         "reviewed": "auto" if tags else "", "tags_src": b.get("tags_src") or ("플랫폼 브랜드 소개문" if tags else "")}
    r["role"] = role(b, sc)
    rows.append(r)

rows.sort(key=lambda r: (-(r["score"].get("score") or -1), -(r["score"].get("partial") or -1), r["ko"]))
for i, r in enumerate(rows, 1):
    r["rank"] = i

# ---------- 브랜드별 파일 ----------
CD = ROOT / "candidates"
CD.mkdir(exist_ok=True)
for old in list(CD.glob("*.md")) + list(CD.glob("*.json")):
    old.unlink()  # 이름이 바뀐 옛 파일이 남지 않게 매번 새로 쓴다
won = lambda v: f"{v:,}원" if isinstance(v, (int, float)) and v is not None else "-"


def pstat(p):
    return " · ".join(f"{k} {p[k]:,}" for k in ("min", "q1", "med", "q3", "p90", "max") if p.get(k) is not None)


def md(r):
    L = [f"# {r['ko']}" + (f" ({r['en']})" if r["en"] else ""), "",
         f"- 상태: **{r['status']}** · 역할: **{r['role']}** · 구분: {r['tier'] or '-'}",
         f"- 종합점수: {r['score'].get('score', '미평가')} (가격위치 {r['score'].get('s_price', '-')} · 가격대겹침 {r['score'].get('s_band', '-')} · 무드 {r['score'].get('s_mood', '-')}) — 가격 출처 {r['price_source'] or '없음'}",
         f"- 무드 태그: {' '.join(r['tags']) or '미부여'} ({r['tags_src'] + '을 읽은 Claude 1차 판단' if r['tags'] else '소개문 미수집'})",
         f"- 찾게 된 경로: {' / '.join(r['found']) or '-'}"]
    if r["note"]:
        L.append(f"- 메모: {r['note']}")
    ch = []
    ids = r["ids"]
    if ids.get("c29"):
        ch.append(f"[29CM](https://www.29cm.co.kr/store/brand/{ids['c29']})")
    if ids.get("wc"):
        ch.append(f"[W컨셉](https://display.wconcept.co.kr/rn/brand/{ids['wc']})")
    if ids.get("ms"):
        ch.append(f"[무신사](https://www.musinsa.com/brand/{ids['ms']}/products?gf=A)")
    if ids.get("off"):
        ch.append(f"[공식/그룹몰]({ids['off']})")
    L += [f"- 채널: {' · '.join(ch) or '미확인'}", ""]
    if r["intro"] or r["sub"]:
        L += ["## 브랜드 소개 (플랫폼 브랜드 정보 원문)", ""]
        if r["sub"]:
            L.append(f"> {r['sub']}")
        if r["intro"]:
            L.append(f"> {re.sub(r'<br ?/?>', ' ', r['intro'])}")
        L.append(f"\n설립연도 {r['est'] or '-'} · 29CM 브랜드 좋아요 {r['heart']:,}" if r["heart"] else f"\n설립연도 {r['est'] or '-'}")
        L.append("")
    for name, snap in (("29CM", r["s29"]), ("W컨셉", r["swc"])):
        if not snap:
            continue
        p = snap["p"]
        d = snap.get("disc") or {}
        L += [f"## 가격·할인 — {name} 브랜드 페이지 ({snap['observed_at']}, 첫 화면 상품 {snap['n']}개)", "",
              f"- 화면 판매가 분포(원): {pstat(p)}",
              f"- 할인 중인 상품 {d.get('n', 0)}개 · 중앙 할인율 {d.get('med') or '-'}% · 최대 {d.get('max') or '-'}%"]
        if name == "29CM":
            L += [f"- 상품 좋아요 중앙 {snap['likes']['med'] if snap['likes']['med'] is not None else '-'} · 표본 리뷰 합 {snap['rev']['total']:,} · 평점 {snap.get('rate') or '-'}", f"- 배지: {json.dumps(snap.get('tags', {}), ensure_ascii=False)} · 쿠폰/혜택: {' / '.join(snap.get('promo', [])) or '-'}"]
        L.append(f"- 샘플 상품: {' | '.join(snap.get('names', [])[:6])}")
        L.append("")
    sc = r["score"]
    if sc:
        L += ["## 정가 환산(점수 계산용)", "", f"- 정가 환산 Q1·중앙·Q3: {won(sc['list_q1'])} · {won(sc['list_med'])} · {won(sc['list_q3'])} (환산배수 {sc['k']})",
              f"- 그레이스유 공식몰 정가 중앙 {won(GP['med'])} (Q1~Q3 {won(GP['q1'])}~{won(GP['q3'])}) 대비 {sc['list_med'] / GP['med']:.2f}배", ""]
    o = r["off"]
    if o:
        L += ["## 공식몰 상의 가격 분포 (공식몰 '상의' 분류 전체, 상품 단위·색 합침)", "",
              f"- 상의 {o['styles']}개 · 판매가 Q1·중앙·Q3 {won(o['sale']['q1'])} · {won(o['sale']['med'])} · {won(o['sale']['q3'])} · 정가 중앙 {won(o['list']['med'])}",
              f"- 표시 할인: 상품의 {round(o['disc_share'] * 100)}% · 중앙 {round((o.get('disc_med') or 0) * 100)}% · 그레이스유 대비 판정: 판매가 {o.get('judge_sale', '-')} / 정가 {o.get('judge_list', '-')}",
              f"- 출처: {o['source']} ({o['observed_at']}) · 상품별 값은 `onepager/상의-상품목록.csv`"]
        L.append("")
    if r["evidence"]:
        L += ["## 찾게 된 근거 — 어패럴뉴스 기사 (이름이 나온 기사를 자동으로 찾은 것이라 동명 단어가 섞일 수 있음)", ""]
        for e in r["evidence"][:8]:
            extra = f" · {e['anchor']}와 같이 언급" if e.get("anchor") else ""
            L.append(f"- {e['date']} [{e['title']}]({e['url']}) ({e['where']}{extra})")
            if e.get("excerpt"):
                L.append(f"  - 발췌: …{e['excerpt']}…")
        L.append("")
    L += ["## 한계·다음에 할 것", "",
          "- 플랫폼 브랜드 페이지는 첫 화면(추천·신상 중심) 표본이라 전체 가격 분포와 다를 수 있다.",
          "- 무드 태그는 위 근거를 읽은 1차 판단이라 사용자 검토 전이다." if r["tags"] else "- 무드 태그가 없어 무드 점수는 비어 있다(소개문을 읽은 뒤 부여).",
          f"- 스냅샷 날짜: 29CM {', '.join(r['snap_days']['29cm']) or '-'} · W컨셉 {', '.join(r['snap_days']['wconcept']) or '-'} "
          f"(`python tools/snap_29cm.py {ids.get('c29', '')}` / `python tools/snap_wconcept.py {ids.get('wc', '')}` 로 갱신, `tools/diff_snapshots.py` 로 변화 비교)", ""]
    return "\n".join(L)


for r in rows:
    (CD / f"{r['slug']}.md").write_text(md(r), encoding="utf-8")
    j = {k: v for k, v in r.items() if k not in ("s29", "swc")}
    j["platform"] = {"29cm": {k: v for k, v in (r["s29"] or {}).items() if k not in ("cards",)} if r["s29"] else None,
                     "wconcept": {k: v for k, v in (r["swc"] or {}).items() if k not in ("cards",)} if r["swc"] else None}
    (CD / f"{r['slug']}.json").write_text(json.dumps(j, ensure_ascii=False, indent=1), encoding="utf-8")

# ---------- INDEX.md ----------
by_role = {}
for r in rows:
    by_role.setdefault(r["role"], []).append(r)
order = ["직접 경쟁(확정)", "가격·무드 모두 근접", "가격 근접·무드 다름", "무드 근접·가격 다름", "가격 근접·무드 미평가", "저가 참고(상의 후보에서 제외)", "가격 앵커(저가)", "가격 앵커(고가)", "참고(트렌드·무드 탐색)", "비의류·참고", "미평가(가격 미수집)"]
IX = [f"# 후보 {len(rows)}곳 목록 (역할별)", "", f"기준: {date.today().isoformat()} · 그레이스유 정가 중앙 {won(GP['med'])} · 점수 산식은 `README.md` §2", ""]
for ro in order:
    if ro not in by_role:
        continue
    IX += [f"## {ro} ({len(by_role[ro])})", "", "| 브랜드 | 상태 | 종합 | 가격 | 무드 | 정가환산 중앙 | 출처 | 파일 |", "|---|---|---|---|---|---|---|---|"]
    for r in by_role[ro]:
        s = r["score"]
        IX.append(f"| {r['ko']} | {r['status']} | {s.get('score', '-')} | {s.get('s_price', '-')} | {s.get('s_mood', '-')} | {won(s['list_med']) if s else '-'} | {r['price_source'] or '-'} | [{r['slug']}.md]({r['slug']}.md) |")
    IX.append("")
(CD / "INDEX.md").write_text("\n".join(IX), encoding="utf-8")

# ---------- 기준(그레이스유) 프로필 ----------
GL = ["# 그레이스유 (GRACE U) — 비교 기준", "",
      "- 공식몰: https://graceu.co.kr (카페24) · W컨셉 브랜드관 108809 · 29CM 99166(현재 상품 없음)",
      "- 공식몰 메타 키워드: 20대후반여성의류, 30대초반여성의류, 원피스, 데일리룩, 커플룩, 셀프웨딩원피스, 셀프웨딩, 베일, 시밀러룩",
      f"- 무드 태그(추정): {' '.join(sorted(GT, key=list(GU['tags']).index))} — 근거: {GU['basis']}", ""]
if bw:
    GL += ["## 브랜드 소개 (W컨셉 브랜드관 원문)", "", f"> {bw.get('intro', '')}", ""]
    GL += [f"## W컨셉 브랜드관 가격 ({bw['observed_at']}, 첫 화면 상품 {bw['n']}개)", "", f"- 화면 판매가 분포(원): {pstat(bw['p'])}",
           f"- 할인 중인 상품 {bw['disc']['n']}개 · 중앙 할인율 {bw['disc']['med']}% · 최대 {bw['disc']['max']}%", ""]
GL += ["## 공식몰 가격 (samples/graceyou-overview/data/products.json, 색 옵션을 합친 상품 기준)", "",
       f"- 전체 {GP['n']}개: 정가 Q1·중앙·Q3 {won(GP['q1'])} · {won(GP['med'])} · {won(GP['q3'])} (P90 {won(GP['p90'])}, 최고 {won(GP['max'])} — 고가 이상치가 있어 평균은 쓰지 않음)",
       f"- 상의(사이트 상의 분류, 색 합침) {GU['price_uniq_top']['n']}개: 판매가 Q1·중앙·Q3 {won(GU['price_uniq_top']['q1'])} · {won(GU['price_uniq_top']['med'])} · {won(GU['price_uniq_top']['q3'])} (프로젝트 기준선)",
       "- 공식몰은 거의 할인하지 않는다(할인 상품 2개). W컨셉에서는 약 절반의 상품에 중앙 13% 할인이 붙는다.", "",
       "## 비교할 때", "", "- 점수의 가격 기준은 공식몰 전체 정가 중앙(213,000원)이다. 플랫폼 화면가(할인 반영)는 정가로 환산해 비교한다.",
       "- 같은 플랫폼끼리 비교하려면 W컨셉 브랜드관 값(위)과 각 후보의 W컨셉 값을 쓴다.", ""]
(CD / "_graceu.md").write_text(chr(10).join(GL), encoding="utf-8")

# ---------- CSV / XLSX ----------
COLS = [("rank", "순위"), ("ko", "브랜드"), ("en", "영문명"), ("status", "상태"), ("role", "역할"), ("tier", "구분"), ("tags", "무드태그"), ("score", "종합점수"), ("s_price", "가격위치"), ("s_band", "가격대겹침"), ("s_mood", "무드"),
        ("partial", "가격만점수"), ("price_source", "가격출처"), ("list_q1", "정가환산Q1"), ("list_med", "정가환산중앙"), ("list_q3", "정가환산Q3"), ("k", "환산배수"), ("p_med", "화면가중앙"), ("disc_n", "할인상품수"), ("disc_med", "중앙할인율"), ("n", "표본수"),
        ("heart", "29CM좋아요"), ("est", "설립연도"), ("sub", "한줄소개"), ("found", "찾은경로"), ("evid", "기사근거수"), ("c29", "29CM"), ("wc", "W컨셉"), ("off", "공식/그룹몰"), ("note", "메모"), ("file", "프로필파일")]


def flat(r):
    s, src = r["score"], (r["s29"] or r["swc"] or {})
    return {"rank": r["rank"], "ko": r["ko"], "en": r["en"], "status": r["status"], "role": r["role"], "tier": r["tier"], "tags": " ".join(r["tags"]),
            "score": s.get("score", ""), "s_price": s.get("s_price", ""), "s_band": s.get("s_band", ""), "s_mood": s.get("s_mood", ""), "partial": s.get("partial", ""),
            "price_source": r["price_source"], "list_q1": s.get("list_q1", ""), "list_med": s.get("list_med", ""), "list_q3": s.get("list_q3", ""), "k": s.get("k", ""),
            "p_med": (src.get("p") or {}).get("med", ""), "disc_n": (src.get("disc") or {}).get("n", ""), "disc_med": (src.get("disc") or {}).get("med", ""), "n": src.get("n", ""),
            "heart": r["heart"] or "", "est": r["est"] or "", "sub": r["sub"], "found": " / ".join(r["found"]), "evid": len(r["evidence"]),
            "c29": r["ids"].get("c29", ""), "wc": r["ids"].get("wc", ""), "off": r["ids"].get("off", ""), "note": r["note"], "file": f"candidates/{r['slug']}.md"}


with open(ROOT / "candidates.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow([h for _, h in COLS])
    for r in rows:
        d = flat(r)
        w.writerow([d[k] for k, _ in COLS])

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "후보풀"
    ws.append([h for _, h in COLS])
    for r in rows:
        d = flat(r)
        ws.append([d[k] for k, _ in COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    wd = {"한줄소개": 44, "찾은경로": 40, "무드태그": 26, "역할": 24, "프로필파일": 34, "공식/그룹몰": 28, "메모": 30}
    for i, (_, h) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd.get(h, 11)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    we = wb.create_sheet("근거기사")
    we.append(["브랜드", "날짜", "기사 제목", "URL", "이름이 나온 곳", "같이 언급된 기준 브랜드", "발췌"])
    for r in rows:
        for e in r["evidence"]:
            we.append([r["ko"], e["date"], e["title"], e["url"], e["where"], e.get("anchor", ""), e.get("excerpt", "")])
    for c in we[1]:
        c.font = Font(bold=True)
    for col, wdt in zip("ABCDEFG", (14, 11, 60, 50, 12, 16, 80)):
        we.column_dimensions[col].width = wdt

    wg = wb.create_sheet("그레이스유 기준")
    wg.append(["항목", "값"])
    for k, v in [("무드태그(추정)", " ".join(GT)), ("기준 근거", GU["basis"]), ("공식몰 가격(색 합친 263개) Q1·중앙·Q3", f"{GP['q1']:,} · {GP['med']:,} · {GP['q3']:,}"), ("공식몰 상의(55개, 색 합침) 판매가 Q1·중앙·Q3", "101,500 · 153,000 · 218,000"),
                 ("W컨셉 브랜드관(첫 화면)", json.dumps(gu_wc, ensure_ascii=False) if gu_wc else "미수집"), ("기준 소개글(다른 에이전트 판정 문서)", "Elegant, timeless, 우아한, 타임리스, lifestyle, 데일리룩 — planning/mood-judgment-v2.md")]:
        wg.append([k, v])
    wg.column_dimensions["A"].width = 38
    wg.column_dimensions["B"].width = 110

    wa = wb.create_sheet("산식·역할")
    for row in [["항목", "설명"], ["가격 위치(0.30)", "정가 환산 중앙값과 그레이스유 중앙값(213,000원)의 로그 거리. 2.5배 차이면 0점"], ["가격대 겹침(0.20)", "정가 환산 Q1~Q3와 그레이스유 Q1~Q3의 로그 구간 겹침 비율"],
                ["무드(0.50)", "소개문에서 붙인 태그와 그레이스유 태그의 자카드 비율. 태그가 없으면 점수를 비운다"], ["정가 환산", "표본 절반 이상이 할인 중이면 화면가 ÷ (1-중앙할인율)"],
                ["역할", "가격 앵커: 환산 중앙이 그레이스유의 0.5배 미만(저가)/1.6배 초과(고가). 가격·무드 모두 근접: 가격점수 70↑ & 무드 30↑. 가격 근접·무드 다름: 가격점수 70↑. 무드 근접·가격 다름: 무드 40↑"],
                ["쓰지 않는 값", "상품 수·옷 카테고리·좋아요·리뷰 수는 점수에 넣지 않는다(맥락으로만 보관)"]]:
        wa.append(row)
    wa.column_dimensions["A"].width = 22
    wa.column_dimensions["B"].width = 120
    wb.save(ROOT / "candidates.xlsx")
    print("xlsx ok")
except ImportError:
    print("openpyxl 없음")

print(len(rows), "후보,", sum(len(r["evidence"]) for r in rows), "근거 연결")
for ro in order:
    if ro in by_role:
        print(ro, len(by_role[ro]), ", ".join(r["ko"] for r in by_role[ro][:12]))
